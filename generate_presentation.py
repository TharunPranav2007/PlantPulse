import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    
    # Set Widescreen 16:9 Aspect Ratio
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6] # Blank Layout

    # Design System Colors
    DARK_BG = RGBColor(15, 23, 42)       # #0f172a (Deep Slate/Navy)
    CARD_BG = RGBColor(30, 41, 59)       # #1e293b (Elevated Slate Card)
    BORDER_COLOR = RGBColor(51, 65, 85) # #334155
    WHITE = RGBColor(255, 255, 255)
    CYAN = RGBColor(0, 242, 254)         # #00f2fe (PlantPulse Cyan Accent)
    BLUE_ACCENT = RGBColor(2, 132, 199)  # #0284c7
    SLATE_TEXT = RGBColor(148, 163, 184) # #94a3b8 (Muted Slate)
    LIGHT_TEXT = RGBColor(203, 213, 225) # #cbd5e1

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_card_shape(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    def add_header(slide, title_text, category_text="U21CS504 - WEB TECHNOLOGIES LABORATORY"):
        # Header Container Shape
        add_card_shape(slide, Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.85), bg_color=CARD_BG, border_color=BORDER_COLOR)
        
        tx_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.42), Inches(11.3), Inches(0.8))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)
        
        p1 = tf.paragraphs[0]
        p1.text = category_text.upper()
        p1.font.name = "Arial"
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = CYAN

        p2 = tf.add_paragraph()
        p2.text = title_text
        p2.font.name = "Arial"
        p2.font.size = Pt(20)
        p2.font.bold = True
        p2.font.color.rgb = WHITE

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (SOLO PROJECT - EXACT DETAILS REQUESTED)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, DARK_BG)

    # Hero Backdrop Card
    add_card_shape(slide1, Inches(1.5), Inches(0.8), Inches(10.333), Inches(5.9), bg_color=CARD_BG, border_color=CYAN)

    tx1 = slide1.shapes.add_textbox(Inches(1.8), Inches(1.1), Inches(9.733), Inches(5.3))
    tf1 = tx1.text_frame
    tf1.word_wrap = True

    # Course Details (No heading label prefixes as requested)
    p = tf1.paragraphs[0]
    p.text = "U21CS504 - WEB TECHNOLOGIES LABORATORY"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = CYAN
    p.alignment = PP_ALIGN.CENTER

    p = tf1.add_paragraph()
    p.text = "PlantPulse"
    p.font.name = "Arial"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(12)

    p = tf1.add_paragraph()
    p.text = "Smart Industrial Asset & Predictive Maintenance Platform"
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = LIGHT_TEXT
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(6)

    p = tf1.add_paragraph()
    p.text = "(Monitor. Maintain. Predict.)"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.italic = True
    p.font.color.rgb = CYAN
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(4)

    # Divider Line Text
    p = tf1.add_paragraph()
    p.text = "____________________________________________________"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.color.rgb = BORDER_COLOR
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(16)

    p = tf1.add_paragraph()
    p.text = "Tharun Pranav T"
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(16)

    p = tf1.add_paragraph()
    p.text = "24CS228"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(4)

    p = tf1.add_paragraph()
    p.text = "Department of Computer Science and Engineering"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.color.rgb = LIGHT_TEXT
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(8)

    p = tf1.add_paragraph()
    p.text = "KPR Institute of Engineering and Technology"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(4)

    # Helper function for text content cards
    def add_content_bullets(slide, left, top, width, height, title, items):
        add_card_shape(slide, left, top, width, height)
        tx = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tx.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Arial"
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = CYAN

        for item in items:
            p = tf.add_paragraph()
            p.text = "• " + item
            p.font.name = "Arial"
            p.font.size = Pt(13)
            p.font.color.rgb = LIGHT_TEXT
            p.space_before = Pt(8)

    # Helper function for slide with left bullets & right screenshot
    def create_screenshot_slide(prs, title, bullets, image_path):
        slide = prs.slides.add_slide(blank_layout)
        set_slide_background(slide, DARK_BG)
        add_header(slide, title)

        # Left Card (Textual Details)
        add_content_bullets(slide, Inches(0.8), Inches(1.45), Inches(4.5), Inches(5.6), "Key Capabilities & Details", bullets)

        # Right Card (Screenshot Image)
        add_card_shape(slide, Inches(5.5), Inches(1.45), Inches(7.033), Inches(5.6), bg_color=CARD_BG, border_color=CYAN)
        
        if os.path.exists(image_path):
            slide.shapes.add_picture(image_path, Inches(5.6), Inches(1.55), Inches(6.833), Inches(5.4))
        return slide

    # =========================================================================
    # SLIDE 2: PROJECT OVERVIEW & PROBLEM STATEMENT
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, DARK_BG)
    add_header(slide2, "1. Project Overview & Industrial Problem Statement")

    add_content_bullets(slide2, Inches(0.8), Inches(1.45), Inches(5.6), Inches(5.6), "Industrial Background", [
        "Industry 4.0 smart manufacturing relies on continuous machinery visibility.",
        "Capital-intensive equipment (5-Axis CNC Mills, Robotic Arms, Stamping Presses) operates under extreme thermal and mechanical stress.",
        "Unplanned breakdowns result in thousands of dollars in lost hourly production.",
        "Conventional plants rely on reactive 'break-fix' repairs or manual paper logbooks."
    ])

    add_content_bullets(slide2, Inches(6.9), Inches(1.45), Inches(5.6), Inches(5.6), "Core Problem Statement", [
        "Manual telemetry recording leads to unnoticed vibration and temperature spikes.",
        "Paper work order dispatches get delayed, misallocated, or lost.",
        "Unmonitored spare parts inventory leads to critical stockouts during repairs.",
        "Existing software lacks role-aware navigation and suffers from severe dark/light theme flickering during user interaction."
    ])

    # =========================================================================
    # SLIDE 3: OBJECTIVES & UNSDG ALIGNMENT
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, DARK_BG)
    add_header(slide3, "2. Project Objectives & UNSDG Alignment")

    add_content_bullets(slide3, Inches(0.8), Inches(1.45), Inches(5.6), Inches(5.6), "Project Objectives", [
        "Develop a multi-role industrial platform supporting 4 distinct roles.",
        "Implement Role-Based Access Control (RBAC) with route security guards.",
        "Engineered a zero-flicker dual-theme design system (Dark & Light) adhering to WCAG 2.1 AAA contrast.",
        "Implement automated predictive telemetry health scoring algorithms.",
        "Build a native PHP 8.x PDO REST API connected to MySQL Server 8.0."
    ])

    add_content_bullets(slide3, Inches(6.9), Inches(1.45), Inches(5.6), Inches(5.6), "UN Sustainable Development Goals", [
        "SDG 9: Industry, Innovation & Infrastructure - Fosters resilient digital manufacturing infrastructure.",
        "SDG 11: Sustainable Cities & Communities - Reduces industrial waste through machine health optimization.",
        "SDG 12: Responsible Consumption & Production - Extends machine lifespans and optimizes spare parts utilization."
    ])

    # =========================================================================
    # SLIDE 4: SYSTEM ARCHITECTURE & DATA FLOW
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, DARK_BG)
    add_header(slide4, "3. 3-Tier System Architecture & Data Flow")

    add_content_bullets(slide4, Inches(0.8), Inches(1.45), Inches(11.733), Inches(5.6), "Architectural Layers & Data Synchronization", [
        "Presentation Layer: Semantic HTML5, Vanilla CSS3 Design Tokens, ES6 JavaScript UI Controllers, and Chart.js 4.x visualizations.",
        "Application API Layer: Central Event Bus (js/store.js), RBAC Guard (js/auth.js), and native PHP 8.x REST API endpoints (api/*.php).",
        "Database Persistence Layer: Relational MySQL Server 8.0 database (plantpulse_db) with 9 normalized tables and foreign key constraints.",
        "Hybrid Auto-Sync Engine: Automatically syncs frontend state with MySQL REST API when live, with seamless fallback to LocalStorage."
    ])

    # =========================================================================
    # SLIDE 5: TECHNOLOGY STACK & ACADEMIC CONSTRAINTS
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, DARK_BG)
    add_header(slide5, "4. Technology Stack & Academic Constraints")

    add_content_bullets(slide5, Inches(0.8), Inches(1.45), Inches(5.6), Inches(5.6), "Stage 1 & Stage 2 Tech Stack", [
        "Structure: Semantic HTML5 Layouts & Validation Modals.",
        "Styling: Vanilla CSS3 Tokens, Flexbox, CSS Grid, Glassmorphism.",
        "Client Scripting: Vanilla ES6 JavaScript (No frameworks).",
        "Backend API: Native PHP 8.x PDO Object-Oriented REST API.",
        "Database: MySQL Server 8.0 (plantpulse_db) with 9 tables."
    ])

    add_content_bullets(slide5, Inches(6.9), Inches(1.45), Inches(5.6), Inches(5.6), "Academic Constraints Enforced", [
        "STRICT COMPLIANCE: React, Angular, Vue, Node.js, Express, MongoDB, Firebase, and PostgreSQL are NOT used.",
        "PDO Prepared Statements: 100% protection against SQL Injection attacks.",
        "Zero-Flicker Early Script: Head script prevents FOUC in Light mode.",
        "Responsive Grid: Optimized for Desktop, Tablet, and Mobile devices."
    ])

    # =========================================================================
    # SLIDES 6 TO 13: MODULE SCREENSHOTS & DETAILED EXPLANATIONS
    # =========================================================================

    # Slide 6: Login Portal
    create_screenshot_slide(prs, "5. Multi-Role Authentication Gateway", [
        "Split-screen desktop gateway with industrial backdrop.",
        "Supports 4 distinct roles: Plant Manager (Admin), Technician, Supervisor, and Inventory Manager.",
        "Features quick-fill demo credentials chips for 1-click evaluation.",
        "Includes live theme toggle, password eye visibility toggle, and loading button spinners."
    ], "screenshots/01_login_portal.png")

    # Slide 7: Admin Dashboard
    create_screenshot_slide(prs, "6. Industrial Command Center Dashboard", [
        "Overview of manufacturing assets, fleet uptime %, active alerts, and open work order counts.",
        "Real-time machine health telemetry overview cards.",
        "Machine Health Trend line chart built with Chart.js.",
        "Real-time activity audit feed with 10-second relative timestamp ticker ('Just now', '15 sec ago')."
    ], "screenshots/02_admin_dashboard.png")

    # Slide 8: Assets Management
    create_screenshot_slide(prs, "7. Industrial Assets Directory & Registration", [
        "Comprehensive machine asset directory with health progress bars and status badges.",
        "Telemetry status indicators tracking vibration (mm/s) and operating temperature (°C).",
        "Interactive '+ Register New Asset' modal form with real-time validation.",
        "Executes POST/PUT/DELETE requests synced directly to MySQL database."
    ], "screenshots/03_assets_management.png")

    # Slide 9: Work Orders Kanban
    create_screenshot_slide(prs, "8. Work Orders Kanban Lifecycle & Queue", [
        "Interactive Kanban board managing work orders across OPEN, ASSIGNED, IN PROGRESS, RESOLVED, and CLOSED status columns.",
        "Technician assignment and supervisor dispatch workflow.",
        "Technicians can mark work started, enter completion notes, and advance status.",
        "Live sync with MySQL work_orders table."
    ], "screenshots/04_workorders_kanban.png")

    # Slide 10: Spare Parts Inventory
    create_screenshot_slide(prs, "9. Spare Parts Inventory & Movement History", [
        "Spare parts catalog tracking stock levels, minimum safety thresholds, and unit costs.",
        "Automated stock status indicators: NORMAL, LOW STOCK, and OUT OF STOCK.",
        "Quick +20 Restock trigger and stock movement log tracking IN/OUT transactions.",
        "Triggers automated low-stock telemetry alerts in alert center."
    ], "screenshots/05_spareparts_inventory.png")

    # Slide 11: Analytics Dashboard
    create_screenshot_slide(prs, "10. Real-Time Analytics & Chart Visualizations", [
        "Interactive Chart.js visual charts: Unplanned Downtime Bar Chart, Cost Distribution Doughnut Chart, and Machine Distribution Polar Area Chart.",
        "Live period, unit, and machine type filter controls.",
        "100% Theme-Adaptive: Grid lines, ticks, tooltips, and legends update dynamically on theme toggle."
    ], "screenshots/06_analytics_dashboard.png")

    # Slide 12: Alert Center & Predictive Engine
    create_screenshot_slide(prs, "11. Predictive Health Scoring & Telemetry Alarms", [
        "Central Alert Center highlighting CRITICAL, WARNING, and INFO telemetry alarms.",
        "Mathematical health score degradation scoring analyzing vibration, temperature, and operating hours.",
        "Interactive formula modal detailing machine degradation factors.",
        "Mark-as-read status toggle synced with MySQL database."
    ], "screenshots/07_alert_center.png")

    # Slide 13: Zero-Flicker Dual Theme
    create_screenshot_slide(prs, "12. Zero-Flicker Dual-Theme Architecture", [
        "Centralized theme system (js/theme.js) managing Dark Slate and Clean Light modes.",
        "Synchronous inline <head> script prevents Dark Mode flash during cross-page navigation or refresh.",
        "High-contrast slate design tokens (--text-section-heading) meeting WCAG 2.1 AAA contrast rules.",
        "Flyout tooltips for collapsed sidebar navigation."
    ], "screenshots/08_light_mode_dashboard.png")

    # =========================================================================
    # SLIDE 14: MYSQL DATABASE SCHEMA & SECURITY
    # =========================================================================
    slide14 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide14, DARK_BG)
    add_header(slide14, "13. Relational MySQL Database Schema (plantpulse_db)")

    add_content_bullets(slide14, Inches(0.8), Inches(1.45), Inches(5.6), Inches(5.6), "9 Relational MySQL Tables", [
        "users: User accounts & RBAC role keys.",
        "assets: Machine records, health & telemetry.",
        "maintenance: Scheduled preventive & corrective PM.",
        "work_orders: Work order lifecycle & assignments.",
        "technicians: Technician directory & workload.",
        "spare_parts: Inventory quantities & costs.",
        "alerts: Telemetry alarms & notifications.",
        "activity_log: Audit trail of all plant actions.",
        "stock_movements: IN/OUT stock transaction history."
    ])

    add_content_bullets(slide14, Inches(6.9), Inches(1.45), Inches(5.6), Inches(5.6), "Security & Data Integrity", [
        "PDO Prepared Statements: Parameterized queries protect against 100% of SQL Injection attacks.",
        "Referential Integrity: Primary Keys & Foreign Keys with CASCADE / SET NULL rules.",
        "Prepared Schema Script: plantpulse_schema.sql handles database setup & seed data.",
        "REST API Layer: api/*.php endpoints format clean JSON responses."
    ])

    # =========================================================================
    # SLIDE 15: CONCLUSION & FUTURE ENHANCEMENTS
    # =========================================================================
    slide15 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide15, DARK_BG)
    add_header(slide15, "14. Conclusion & Future Roadmap")

    add_content_bullets(slide15, Inches(0.8), Inches(1.45), Inches(5.6), Inches(5.6), "Project Achievements", [
        "Delivered full Stage 1 (Static Web App) & Stage 2 (PHP + MySQL Backend) requirements.",
        "Eliminated paper work order dispatch delays and reduced maintenance response times.",
        "Achieved zero-flicker dual-theme stability and high-contrast accessibility.",
        "Demonstrated real-time data flow between browser UI and MySQL Server 8.0."
    ])

    add_content_bullets(slide15, Inches(6.9), Inches(1.45), Inches(5.6), Inches(5.6), "Future Scalability Roadmap", [
        "IoT Sensor Integration: Connect physical MQTT vibration and thermal sensors for live hardware feeds.",
        "AI Deep Learning Models: Implement Python TensorFlow predictive failure forecasting.",
        "Mobile App Companion: Build a mobile PWA companion for field maintenance technicians."
    ])

    output_filename = "PlantPulse_Web_Technologies_Presentation.pptx"
    prs.save(output_filename)
    print(f"Presentation successfully generated: {output_filename}")

if __name__ == "__main__":
    create_presentation()
